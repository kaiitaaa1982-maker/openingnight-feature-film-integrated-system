"""Read generated files using independent Excel/PDF readers; no workbook rewrite."""
import json,sys,unicodedata
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
import pypdfium2 as pdfium
out=Path(sys.argv[1] if len(sys.argv)>1 else 'data/report-completion-20260923')
files={}
for path in out.glob('*.xlsx'):
    book=load_workbook(path,data_only=False,read_only=True)
    rows={sheet.title:list(sheet.values) for sheet in book}
    assert not any(cell.data_type=='f' for sheet in book for row in sheet for cell in row)
    if path.stem=='committee-income':
        pool=next(row for row in rows['収支報告'] if row[0]=='分配原資')
        assert pool[1:]==(800000,1600000,2400000),pool
        assert [r[3] for r in rows['出資者分配'][1:]]==[800000,480000,320000]
    files[path.name]={'sheets':list(rows),'rows':{k:len(v) for k,v in rows.items()}}
    book.close()
for path in out.glob('*.pdf'):
    reader=PdfReader(path);texts=[p.extract_text() for p in reader.pages]
    assert texts and all(len(t)>50 for t in texts)
    if path.stem=='committee-income':
        combined=unicodedata.normalize('NFKC','\n'.join(texts))
        for heading in ['収支報告','出資者分配','製作費と出資','損益収支経緯']:
            assert heading in combined,heading
        assert '1600000' in texts[0].replace(',','')
    doc=pdfium.PdfDocument(path)
    doc[0].render(scale=1.2).to_pil().save(out/(path.stem+'-pdf-page1.png'))
    files[path.name]={'pages':len(reader.pages),'textExtracted':True}
(out/'file-verification.json').write_text(json.dumps(files,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(files,ensure_ascii=False))
