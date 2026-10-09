# 生成元（2026-09-26）。出力した2ブックは代表に渡した版で、配布物は fixtures/xlsx-source/demo-sales-*/ の部品から組み立てる
# （scripts/prepare-test-assets.mjs）。作り直すと作成日時などの部品が変わるので、部品と parts.json は差し替えない。
# 使い方: python generate_demo_sales_samples.py <出力フォルダ>
# 架空組織 DEMO-SALES の架空データ（seed-sales-demo.mjs）に合わせた売上報告サンプル（配信・ビデオグラム）を作る。
# 取引先・商品コード（DEMO-Wxx-DIG / DEMO-Wxx-RNT）は DEMO-SALES のマスタにある値だけを使う。対象月は seed の最終月（2026-06）の翌月。
# 各ファイルに「要修正」の行を1行ずつ入れる（取込の練習用。登録するときは理由を書いて除く）。
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

NOTICE = '非公式・架空データ・実在企業発行書式ではありません。取込・加工テスト専用。'
BOLD = Font(bold=True)
HEAD_FILL = PatternFill('solid', fgColor='E8EEF4')
WARN_FILL = PatternFill('solid', fgColor='FCE8E6')
YEN = '#,##0'

# 配信（都度課金）。PF手数料は TVOD 30%・EST 20%
STREAM_ROWS = [
    # 作品CD, 商品コード, 区分, 件数, 単価, 手数料率, 備考
    ('DEMO-W01', 'DEMO-W01-DIG', 'TVOD', 820, 440, 0.30, '正常'),
    ('DEMO-W01', 'DEMO-W01-DIG', 'EST', 150, 2000, 0.20, '正常'),
    ('DEMO-W04', 'DEMO-W04-DIG', 'TVOD', 610, 440, 0.30, '正常'),
    ('DEMO-W04', 'DEMO-W04-DIG', 'EST', 95, 2000, 0.20, '正常'),
    ('DEMO-W12', 'DEMO-W12-DIG', 'TVOD', 540, 400, 0.30, '正常'),
    ('DEMO-W16', 'DEMO-W16-DIG', 'TVOD', 380, 400, 0.30, '正常'),
    ('DEMO-W19', 'DEMO-W19-DIG', 'TVOD', 1240, 500, 0.30, '正常'),
    ('DEMO-W19', 'DEMO-W19-DIG', 'EST', 210, 2500, 0.20, '正常'),
    ('DEMO-W20', 'DEMO-W20-DIG', 'TVOD', 980, 500, 0.30, '正常'),
]
STREAM_BAD = ('DEMO-W07', '', 'TVOD', 60, 400, 7200, 17000, '要修正: 商品コード欠落・正味額不一致')

# ビデオグラム（レンタル）。正味数量＝出荷−返品、正味額＝正味数量×精算単価
RENTAL_ROWS = [
    ('DEMO-W03', 'DEMO-W03-RNT', 'レンタルDVD', 260, 12, 420, '正常'),
    ('DEMO-W03', 'DEMO-W03-RNT', 'レンタルBD', 90, 4, 580, '正常'),
    ('DEMO-W05', 'DEMO-W05-RNT', 'レンタルDVD', 140, 9, 420, '正常'),
    ('DEMO-W09', 'DEMO-W09-RNT', 'レンタルDVD', 310, 15, 420, '正常'),
    ('DEMO-W09', 'DEMO-W09-RNT', 'レンタルBD', 120, 6, 580, '正常'),
    ('DEMO-W14', 'DEMO-W14-RNT', 'レンタルDVD', 420, 18, 420, '正常'),
    ('DEMO-W17', 'DEMO-W17-RNT', 'レンタルDVD', 180, 7, 420, '正常'),
]
RENTAL_BAD = ('DEMO-W09', 'DEMO-W09-RNT', 'レンタルDVD', 40, 75, 500, 21000, '要修正: 返品超過・正味額不一致')


def header_block(ws, title, period_line, headers):
    for row, text in ((1, title), (2, NOTICE), (3, period_line)):
        ws.cell(row=row, column=1, value=text).font = BOLD if row < 3 else Font()
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers))
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=text)
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal='center')
    ws.freeze_panes = 'A5'
    for col, width in zip('ABCDEFGHIJ', (14, 16, 12, 12, 13, 12, 13, 13, 13, 34)):
        ws.column_dimensions[col].width = width


def streaming(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Details'
    headers = ['作品CD', '商品ｺｰﾄﾞ', '配信区分', '対象月', '視聴/契約数', '単価(税抜)', '総額', 'PF手数料', '正味売上', '備考']
    header_block(ws, '配信プラットフォーム風 売上報告（架空）', '対象期間: 2026-07 / 取引先: 架空配信B・都度課金（架空）/ 6作品', headers)
    gross = fee = net = 0
    r = 5
    for work, sku, kind, qty, price, rate, note in STREAM_ROWS:
        g = qty * price
        f = round(g * rate)
        n = g - f
        gross, fee, net = gross + g, fee + f, net + n
        ws.append([work, sku, kind, '2026/07', qty, price, g, f, n, note])
        r += 1
    work, sku, kind, qty, price, f, n, note = STREAM_BAD
    ws.append([work, sku or None, kind, '2026/07', qty, price, qty * price, f, n, note])
    bad_row = r
    for row in ws.iter_rows(min_row=5, max_row=bad_row):
        for cell in row[6:9]:
            cell.number_format = YEN
    for cell in ws[bad_row]:
        cell.fill = WARN_FILL
    check = wb.create_sheet('検証期待値')
    check.append(['検証項目', '期待値', None, 'Details の行', '総額再計算', '手数料（報告）', '正味再計算', '報告の正味との差'])
    items = [('注意', NOTICE), ('取引先', '架空配信B・都度課金（架空）（DEMO-PF-B）'), ('報告の種類', '配信'), ('正常行数', len(STREAM_ROWS)),
             ('作品数', len({row[0] for row in STREAM_ROWS})), ('総額合計', gross), ('手数料合計', fee), ('正味売上合計', net), ('要修正行', f'Details {bad_row}行目')]
    for i in range(max(len(items), bad_row - 4)):
        row = list(items[i]) if i < len(items) else [None, None]
        src = 5 + i
        if src <= bad_row:
            row += [None, src, f'=Details!E{src}*Details!F{src}', f'=Details!H{src}', f'=E{i + 2}-F{i + 2}', f'=G{i + 2}-Details!I{src}']
        check.append(row)
    for col, width in zip('ABCDEFGH', (16, 44, 2, 12, 14, 14, 14, 16)):
        check.column_dimensions[col].width = width
    for cell in check[1]:
        cell.font = BOLD
    for row in check.iter_rows(min_row=2):
        for cell in row[4:8]:
            cell.number_format = YEN
    check['B7'].number_format = YEN
    check['B8'].number_format = YEN
    check['B9'].number_format = YEN
    wb.save(path)
    return {'rows': len(STREAM_ROWS), 'works': len({row[0] for row in STREAM_ROWS}), 'gross': gross, 'fee': fee, 'net': net, 'badRow': bad_row}


def rental(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Details'
    headers = ['作品コード', '品番', '商材', '集計年月', '出荷数量', '返品数', '正味数量', '精算単価', '正味額', '検証メモ']
    header_block(ws, 'ビデオグラム事業者風 出荷・返品報告（架空）', '対象期間: 2026-07 / 取引先: 架空レンタルチェーン（架空）/ 5作品', headers)
    units = amount = 0
    r = 5
    for work, sku, item, shipped, returned, price, note in RENTAL_ROWS:
        n = shipped - returned
        units, amount = units + n, amount + n * price
        ws.append([work, sku, item, '2026年7月', shipped, returned, n, price, n * price, note])
        r += 1
    work, sku, item, shipped, returned, price, a, note = RENTAL_BAD
    ws.append([work, sku, item, '2026年7月', shipped, returned, shipped - returned, price, a, note])
    bad_row = r
    for row in ws.iter_rows(min_row=5, max_row=bad_row):
        for cell in row[4:9]:
            cell.number_format = YEN
    for cell in ws[bad_row]:
        cell.fill = WARN_FILL
    check = wb.create_sheet('検証期待値')
    check.append(['検証項目', '期待値', None, 'Details の行', '正味数量再計算', '正味額再計算', '報告の正味額との差'])
    items = [('注意', NOTICE), ('取引先', '架空レンタルチェーン（架空）（DEMO-RNT-A）'), ('報告の種類', 'ビデオグラム'), ('正常行数', len(RENTAL_ROWS)),
             ('作品数', len({row[0] for row in RENTAL_ROWS})), ('正味数量合計', units), ('正味額合計', amount), ('要修正行', f'Details {bad_row}行目')]
    for i in range(max(len(items), bad_row - 4)):
        row = list(items[i]) if i < len(items) else [None, None]
        src = 5 + i
        if src <= bad_row:
            row += [None, src, f'=Details!E{src}-Details!F{src}', f'=E{i + 2}*Details!H{src}', f'=F{i + 2}-Details!I{src}']
        check.append(row)
    for col, width in zip('ABCDEFG', (16, 44, 2, 12, 16, 14, 18)):
        check.column_dimensions[col].width = width
    for cell in check[1]:
        cell.font = BOLD
    for row in check.iter_rows(min_row=2):
        for cell in row[4:7]:
            cell.number_format = YEN
    check['B7'].number_format = YEN
    check['B8'].number_format = YEN
    wb.save(path)
    return {'rows': len(RENTAL_ROWS), 'works': len({row[0] for row in RENTAL_ROWS}), 'units': units, 'amount': amount, 'badRow': bad_row}


if __name__ == '__main__':
    import json
    out = sys.argv[1]
    result = {'streaming': streaming(f'{out}/demo-sales-streaming-sample.xlsx'), 'rental': rental(f'{out}/demo-sales-videogram-sample.xlsx')}
    print(json.dumps(result, ensure_ascii=False))
