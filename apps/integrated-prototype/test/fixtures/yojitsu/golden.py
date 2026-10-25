#!/usr/bin/env python3
"""案件収支の予実の golden（架空版。expected.json）を作る。

会社ルート plans/2026-10-08-phase1-five-second-yojitsu.md の香盤表 #23c。#11・#13・#14・#14b・#20 の試験が読む。

- (1) 予実・行動・段階・月別: 種（seed.json）の入力の表を、会社ルート ORG/事業計画/_build/build_yojitsu.py（#23b）の
  --seed・--out で一時フォルダに流し、その expected.json（Excel の式と同じ規則の値。セルの番地で持つ）を、行の鍵・期間・月で
  引ける形に直す。式は build_yojitsu.py のコードそのもので、値だけ架空。--snapshot は渡さない（試算を写さない）
- build_yojitsu.py の稼ぎ方（EARN）が種の5つ（D-039 の 2026-10-10 の追記）と違う間は、読み込んだモジュールの稼ぎ方の語
  （EARN・EARN_SALES・TARGET_EARN と、稼ぎ方の数で決まる行動のシートの行の位置）だけを種の vocabulary に差し替えて流す。
  ファイルは書き換えない。差し替えたかどうかは expected.json の build_yojitsu に書く
- (2) 施策の送信後の期待値: #20 で施策の入力値を固定してから作る。いまは空の枠と理由だけ
- (3) 利益額（K-16）: 売上（入金×売上率）− 案件の経費を、種と入金の確認の請求（receipt-matching/invoices.json）から
  build_yojitsu.py を使わずに計算する。計上と入金の両方。費目を 確定・未確定・除外 に分け、照合するのは確定だけ
- 値の無い経路: 種に行が無く golden では確かめられない費目・案件不明・稼ぎ方の実績を、(1)(3) の値から数えて書く
  （#23e で4区分の行を足したあとに作り直す）
- 同じ種から2回作って同じバイト列。ORG の予実管理表.xlsx・expected.json・trials_snapshot.json は開かない（実行の前後で
  大きさと更新時刻が変わらないことを確かめる）。書く前に、golden に ORG の案件コード・名前・金額が無いことを確かめる

使い方: python golden.py [--build-yojitsu <build_yojitsu.py>] [--out <フォルダ>] [--keep-build <フォルダ>] [--check]
  --build-yojitsu を省くと、環境変数 ON_BUILD_YOJITSU、無ければこのファイルから上へたどって ORG/事業計画/_build/build_yojitsu.py
  --keep-build は build_yojitsu.py の出力（予実管理表.xlsx と expected.json）を残すフォルダ（verify_yojitsu.ps1 -Out で Excel と照らす）。
    この照合は、セルの値が全部一致しても終了コード 1 になる: 種の金額が ORG の桁より大きく、列幅に収まらないセル（月別の出金計の
    予測など）が #### 表示になり failures に数えられる（2026-10-10 に 6,999 セルの値の一致と #### 表示 2 件を確かめた）。
    値の正しさはセルの値の一致で見る。#### 表示は golden の値に影響しない
  --check は一時フォルダに作り、置いてある expected.json と同じバイト列かを確かめる（違えば終了コード 1）
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import os
import re
import shutil
import sys
import tempfile
from fractions import Fraction
from pathlib import Path

sys.dont_write_bytecode = True  # 会社ルートの _build/ に __pycache__ を作らない

HERE = Path(__file__).resolve().parent
RM_DIR = HERE.parent / "receipt-matching"
BUILD_REL = Path("ORG") / "事業計画" / "_build" / "build_yojitsu.py"
NOTICE = "架空データ。案件収支の予実の試験とデモ専用で、実在の案件・作品・取引先・金額・目標ではない"
PLAN = "会社ルート plans/2026-10-08-phase1-five-second-yojitsu.md の #23c"
TOTAL = "5年通算"

# ---------------------------------------------------------------- 利益額（K-16）の費目
# 費目 → (向き, 状態, 理由)。向き: 売上は額に売上率を掛けて足し、経費は引く。状態: 確定だけを照合する
ITEMS = {
    "売上": ("売上", "確定", "K-16: 入金（計上は計上の売上）×案件の売上率"),
    "売上（源泉のある取引先）": ("売上", "未確定", "U-5 の残り（源泉と税込）。外国で源泉される契約の取引先の売上"),
    "売上（作品の無い案件）": ("売上", "未確定", "#14b（U-9・U-5）。作品の無い案件の入金"),
    "自社出資の出資額": ("経費", "確定", "J-11b (a): 払った月に引く。出資の記録の払込（払込が種の時点より後のものは入れない）"),
    "製作委員会の出資額（自社の分）": ("経費", "確定", "J-11b (a): 払った月に引く。出資の記録の自社の払込"),
    "買付の MG": ("経費", "確定", "J-11b (a): 払った月に引く。権利元の取り分から回収した分は戻す（種には回収が無い）"),
    "買い切りの買付額": ("経費", "確定", "J-11b (a)・J-9 (c): 払った月に引く"),
    "宣伝費（負担者 自社）": ("経費", "確定", "K-16: 案件の経費（P&A。負担者が自社）"),
    "配給費（負担者 自社）": ("経費", "未確定", "U-4（配給費の会計の科目）。#24 は配給費の科目を U-4 まで「未確定」とするので、負担者が自社の額も別の行に出す"),
    "宣伝費（負担者 自社以外）": ("経費", "未確定", "U-4 R-d・R-f（#24）。負担者が配給会社・委員会・出資の P&A を引くか"),
    "配給費（負担者 自社以外）": ("経費", "未確定", "U-4 R-d・R-f（#24）。負担者が配給会社・委員会・出資の P&A を引くか"),
    "費用区分「制作費」の経費": ("経費", "未確定", "U-5 の残り（費用区分「制作費」の経費）"),
    "制作の実費": ("経費", "除外", "J-14 (a): 受託は受託費の純額で、制作の実費は粗利に入れない"),
    "出資の実績の行（出資の記録の払込と同じ）": ("経費", "除外", "J-11b (a): 出資額は1回だけ引く。出資の記録の払込で引いている"),
}
ITEM_ORDER = list(ITEMS)
RULES = [
    "利益額＝売上 − 案件の経費（K-16。D-039 の 2026-10-09 の追記2）。税別。給与は引かない",
    "入金の基準（cash）: 売上は入金の行（実績。日付の月）の税別の額×案件の売上率。経費は払った日の月（未払の経費は入れない）",
    "計上の基準（accrual）: 売上は請求明細（sale_lines の写し）の計上月の税別の額×案件の売上率。経費は計上月。"
    "計画の「計上と入金」の表の列で取った仮の取り方で、U-5 の残り（計上側の取り方）が決まったら作り直す",
    "J-11b (a) の4費目（自社出資の出資額・製作委員会の出資額・買付の MG・買い切りの買付額）は、計上・入金とも払った月に引く",
    "売上率を掛けた額は、基準×月×案件×費目ごとに税別の額を足してから1回だけ 0 に向かって切り捨てる（signedRateAmount と同じ）",
    "入金の日付の軸（「入金日」）と入金の税抜の出し方は U-5 の残り。種の入金は請求を全額消した行だけで、税別は請求明細の税抜額",
    "確認が要る入金（pending_receipts）は承認するまで登録しないので入れない。返品の行（return_lines）は計上月が無く、入れない",
    "照合するのは状態が「確定」の費目だけ。「未確定」は U-5・U-4 の答えが出るまで別の行に出す額、「除外」は引かないと決まった額",
]

EMPTY_REASON = ("種に行が無く、golden では値を確かめられない経路。#23e で4区分の行を足したあとに golden を作り直す。"
                "それまで #14b の受け入れ基準のうち『案件不明』の行と作品の無い案件の入金の「未確定」、J-9 (c)（買い切り）・"
                "J-14 (a)（受託）の経路は別の fixture で確かめ、#14b の PR にそう書く")

# ---------------------------------------------------------------- 小さな道具
def lf_bytes(path: Path) -> bytes:
    return path.read_bytes().replace(b"\r\n", b"\n")


def sha256(path: Path) -> str:
    return hashlib.sha256(lf_bytes(path)).hexdigest()


def ym(iso: str) -> str:
    return iso[:7]


def norm(v):
    """Excel の空（期間の外・分母が0）は null。「未設定」はそのまま。"""
    return None if v == "" else v


def keystr(k) -> str:
    return "/".join(k) if isinstance(k, tuple) else k


def rate_amount(amount: int, bps: int) -> int:
    """税別の額×売上率（bps）。0 に向かって切り捨てる（統合システムの signedRateAmount と同じ）。"""
    sign = -1 if amount < 0 else 1
    return sign * (abs(amount) * bps // 10000)


# ---------------------------------------------------------------- build_yojitsu.py を読む
def find_build_yojitsu(arg: str | None) -> Path:
    if arg:
        return Path(arg).resolve()
    if os.environ.get("ON_BUILD_YOJITSU"):
        return Path(os.environ["ON_BUILD_YOJITSU"]).resolve()
    for parent in HERE.parents:
        if (parent / BUILD_REL).is_file():
            return parent / BUILD_REL
    raise SystemExit("build_yojitsu.py が見つからない。--build-yojitsu か環境変数 ON_BUILD_YOJITSU で会社ルートのファイルを渡す")


def load_build_yojitsu(path: Path):
    if not path.is_file():
        raise SystemExit(f"build_yojitsu.py が無い: {path}")
    spec = importlib.util.spec_from_file_location("build_yojitsu", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    for name in ("read_seed", "build", "expected", "main", "COLS", "EARN", "EARN_SALES", "MONTHS", "PERIOD_NAMES", "STEPS", "STAGES"):
        if not hasattr(module, name):
            raise SystemExit(f"build_yojitsu.py に {name} が無い（#23b の引数化の版を渡す）")
    return module


def protected_files(bm) -> list[Path]:
    """build_yojitsu.py が引数なしで読み書きする ORG のファイル。開かず、大きさと更新時刻だけを見る。"""
    return [Path(bm.OUT), Path(bm.EXPECTED), Path(bm.SNAPSHOT)]


def stat_of(paths):
    out = {}
    for p in paths:
        st = p.stat() if p.exists() else None
        out[str(p)] = (st.st_size, st.st_mtime_ns) if st else None
    return out


def act_layout(top_first: int, n_earn: int, n_steps: int, n_periods: int):
    """build_yojitsu.py の行動のシートの行の位置（ACT_LONG_HEAD・ACT_LONG_FIRST・ACT_LONG_LAST・ACT_SUM_HEAD）の決め方。"""
    head = top_first + n_earn + 2
    last = head + 1 + n_earn * n_steps * n_periods - 1
    return head, head + 1, last, last + 3


def use_seed_earnings(bm, vocabulary):
    """build_yojitsu.py の稼ぎ方を種の5つにそろえる。戻りは expected.json の build_yojitsu に書く記録。"""
    types = vocabulary["earning_types"]
    names = [t["name"] for t in types]
    sales = {t["name"]: (t["rate_basis"], t["default_sales_rate_bps"] / 10000, t["rate_status"]) for t in types}
    if list(bm.EARN) == names:
        bad = [e for e in names if bm.EARN_SALES[e][1] != sales[e][1]]
        if bad:
            raise SystemExit(f"build_yojitsu.py の既定の売上率が種と違う稼ぎ方: {bad}")
        return {"稼ぎ方の差し替え": False}
    got = (bm.ACT_LONG_HEAD, bm.ACT_LONG_FIRST, bm.ACT_LONG_LAST, bm.ACT_SUM_HEAD)
    if got != act_layout(bm.ACT_TOP_FIRST, len(bm.EARN), len(bm.STEPS), len(bm.PERIOD_NAMES)) \
            or list(bm.TARGET_EARN) != ["全体", "制作の取り分"] + list(bm.EARN):
        raise SystemExit("build_yojitsu.py の行動のシートの行の位置か目標の稼ぎ方の決め方が変わった。use_seed_earnings を直す")
    before = len(bm.EARN)
    bm.EARN = names
    bm.EARN_SALES = sales
    bm.TARGET_EARN = ["全体", "制作の取り分"] + names
    bm.ACT_LONG_HEAD, bm.ACT_LONG_FIRST, bm.ACT_LONG_LAST, bm.ACT_SUM_HEAD = act_layout(
        bm.ACT_TOP_FIRST, len(names), len(bm.STEPS), len(bm.PERIOD_NAMES))
    return {"稼ぎ方の差し替え": True,
            "理由": f"build_yojitsu.py の稼ぎ方は{before}つ（2026-10-08）のままで、種の5つ（D-039 の 2026-10-10 の追記）を読めない。"
                    "読み込んだモジュールの語だけを差し替えた（ファイルは書き換えない）。build_yojitsu.py が5つになったら差し替えは要らない",
            "差し替えた名前": ["EARN", "EARN_SALES", "TARGET_EARN", "ACT_LONG_HEAD", "ACT_LONG_FIRST", "ACT_LONG_LAST", "ACT_SUM_HEAD"],
            "既定の売上率": "種の vocabulary.earning_types の default_sales_rate_bps"}


def build_input(bm, seed):
    """build_yojitsu.py --seed に渡す写し（入力の表を列の名前だけに絞る。generate.py --build-yojitsu-input と同じ形）。"""
    heads = [bm.ACTION_HEAD[0]] + list(bm.ACTION_KEYS.values())
    out = {name: [{c: row[c] for c in cols} for row in seed[name]] for name, cols in bm.COLS.items()}
    out["行動"] = [{h: row[h] for h in heads} for row in seed["行動"]]
    out["設定"] = seed["設定"]
    out["試算の案件"] = seed["試算の案件"]
    return out


def run_build(bm, seed, work: Path):
    """#23b の走らせ方（--seed と --out。--snapshot なし）で流し、出た expected.json と、セルの番地の割り付けを返す。"""
    inp = work / "build-yojitsu-input.json"
    inp.write_bytes(json.dumps(build_input(bm, seed), ensure_ascii=False, indent=1).encode("utf-8"))
    out = work / "build"
    bm.main(["--seed", str(inp), "--out", str(out)])
    produced = json.loads((out / "expected.json").read_text(encoding="utf-8"))
    data, _ = bm.read_seed(inp)
    wb, g_main, g_month, stage_layout = bm.build(data)
    again = json.loads(json.dumps(bm.expected(data, g_main, g_month, stage_layout), ensure_ascii=False, default=str))
    if again != produced:
        raise SystemExit("build_yojitsu.py の expected.json と、同じ種で組み直した割り付けの期待値が違う")
    return produced, data, wb, g_main, g_month, stage_layout, out


# ---------------------------------------------------------------- (1) セルの番地 → 行の鍵・期間・月
def target_derived(specs, base):
    """目標・計画から出る行（target の行と、それを使う diff・ratio の行）の鍵。"""
    derived = set(base)
    for key, kind, args in specs:
        if kind == "target":
            derived.add(key)
        elif kind in ("diff", "ratio") and (args[0] in derived or args[1] in derived):
            derived.add(key)
    return derived


def grid_rows(bm, ws, g, cells, derived):
    from openpyxl.utils import get_column_letter
    key_of = {r: k for k, r in g.rows.items()}
    spec = {k: (kind, args) for k, kind, args in g.specs}
    rows, section = [], None
    for r in range(g.hr + 2, max(g.rows.values()) + 1):
        if r not in key_of:
            label = ws.cell(r, 1).value
            if isinstance(label, str) and label:
                section = label
            continue
        k = key_of[r]
        kind, args = spec[k]
        periods = {bm.PERIOD_NAMES[p]: norm(cells[f"{get_column_letter(col)}{r}"]) for p, col in enumerate(g.period_cols)}
        periods[TOTAL] = norm(cells[f"{get_column_letter(g.total_col)}{r}"])
        row = {"鍵": keystr(k), "段": section, "項目": ws.cell(r, 1).value, "中身": ws.cell(r, 2).value, "種類": kind}
        if kind in ("diff", "ratio"):
            row["a"], row["b"] = keystr(args[0]), keystr(args[1])
        if k in derived:
            row["fictional"] = True
        row["期間"] = periods
        row["月"] = [norm(cells[f"{get_column_letter(g.first_m + i)}{r}"]) for i in range(len(bm.MONTHS))]
        rows.append(row)
    return rows


def action_section(bm, data, cells):
    from openpyxl.utils import get_column_letter
    top = []
    for e in bm.EARN:
        r, v = bm.act_top_row(e), data["行動"][e]
        top.append({"稼ぎ方": e, "1件あたりの額": v["額"], "売上率": cells[f"C{r}"], "1件の売上": norm(cells[f"D{r}"]),
                    "提案→契約の率": v["提案→契約"], "面談→提案の率": v["面談→提案"], "状態": v["状態"]})
    long = []
    for e in bm.EARN:
        for s in bm.STEPS:
            for lab in bm.PERIOD_NAMES:
                r = bm.act_long_row(e, s, lab)
                long.append({"稼ぎ方": e, "ステップ": s, "期間": lab, "計画の額": norm(cells[f"D{r}"]), "目標件数": norm(cells[f"E{r}"]),
                             "実績件数": cells[f"F{r}"], "達成率": norm(cells[f"G{r}"]), "fictional": True})
    sums = []
    r0 = bm.ACT_SUM_HEAD + 1
    for i, (label, _) in enumerate(bm.ACT_SUM_ROWS):
        sums.append({"項目": label, "fictional": True,
                     "期間": {lab: norm(cells[f"{get_column_letter(2 + p)}{r0 + i}"]) for p, lab in enumerate(bm.PERIOD_NAMES)}})
    return {"上の表": top, "目標件数と実績件数": long, "期間ごとの計画と目標": sums}


def stage_section(bm, ws, cells, layout):
    from openpyxl.utils import get_column_letter
    names3 = ["件数", "案件の売上見込", "1年目の予測"]
    tc = 3 + len(bm.EARN) * 3

    def triple(r, c0):
        return {n: cells[f"{get_column_letter(c0 + k)}{r}"] for k, n in enumerate(names3)}

    rows = []
    for s in bm.STAGES:
        r = layout[s[1]]
        rows.append({"段階": s[1], "確度": s[2], "稼ぎ方": {e: triple(r, 3 + j * 3) for j, e in enumerate(bm.EARN)},
                     "合計": triple(r, tc)})
    bottom = []
    for r, kind, per_earn, _ in layout["下の表"]:
        label = ws.cell(r, 1).value
        row = {"項目": label, "種類": kind}
        if label not in ("着地予測", "実績（売上）"):
            row["fictional"] = True
        if per_earn:
            row["稼ぎ方"] = {e: norm(cells[f"{get_column_letter(3 + j)}{r}"]) for j, e in enumerate(bm.EARN)}
        row["合計"] = norm(cells[f"{get_column_letter(3 + len(bm.EARN))}{r}"])
        bottom.append(row)
    # 表の「合計」の行（列の SUM）は build_yojitsu.py の期待値に無いので持たない（段階ごとの行の和）
    return {"段階ごと": rows, "1年目の計画と進み": bottom}


def flags_section(produced, data):
    out = {}
    for sheet, col, label in (("目標", "L", "目標の重複"), ("計画", "L", "計画の重複"), ("予測", "X", "予測の案件不明"),
                              ("実績", "Q", "実績の案件不明")):
        cells = produced["sheets"][sheet]
        out[label] = [row[0] for i, row in enumerate(data[sheet]) if cells.get(f"{col}{i + 2}")]
    return out


def targets_section(seed):
    """目標・計画の行（#11 が照合する）。月の額は年÷12（単位が毎月なら年の額のまま）。
    持ち主（owner_user_email）は種のまま写す（行動指標の計画 #15: A-4 の利用者 demo-admin@openingnight.invalid）。"""
    for t in seed["目標"] + seed["計画"]:
        if t.get("fictional") is not True or t.get("owner_user_email") != seed["owner_user_email"]:
            raise SystemExit(f"種の目標・計画の行に fictional: true か種の持ち主が無い: {t.get('目標ID') or t.get('計画ID')}")
    targets = [{"目標ID": t["目標ID"], "版": t["版"], "目標の種類": t["目標の種類"], "期間": t["期間"], "稼ぎ方": t["稼ぎ方"],
                "単位": t["単位"], "年の額": t["目標額（円）"],
                "月の額": t["目標額（円）"] / 12 if t["単位"] == "期間の合計" else t["目標額（円）"],
                "状態": t["状態"], "owner_user_email": t["owner_user_email"], "fictional": True} for t in seed["目標"]]
    plans = [{"計画ID": p["計画ID"], "版": p["版"], "期間": p["期間"], "稼ぎ方": p["稼ぎ方"], "年の額": p["計画額（円・税別）"],
              "月の額": p["計画額（円・税別）"] / 12, "状態": p["状態"], "owner_user_email": p["owner_user_email"],
              "fictional": True} for p in seed["計画"]]
    return {"目標": targets, "計画": plans}


def check_targets_against_build(bm, seed, rows):
    """目標・計画の月の額（年÷12）が、build_yojitsu.py の予実の月の値と同じかを確かめる。"""
    by_key = {r["鍵"]: r for r in rows}
    periods = [(p["期間"], p["始め"][:7], p["終わり"][:7]) for p in seed["設定"]["期間"]]
    months = [m.isoformat()[:7] for m in bm.MONTHS]

    def label(m):
        return next((n for n, s, e in periods if s <= m <= e), "")

    for i, m in enumerate(months):
        lab = label(m)
        if not lab:
            continue
        for key, kind in (("目標本人", "本人"), ("目標会社", "会社")):
            rows_ = [t["目標額（円）"] for t in seed["目標"] if t["目標の種類"] == kind and t["期間"] == lab and t["稼ぎ方"] == "全体"
                     and t["単位"] == "期間の合計" and t["状態"] == "有効"]
            want = sum(rows_) / 12 if rows_ else "未設定"
            if by_key[key]["月"][i] != want:
                raise SystemExit(f"{key} の {m} が目標の行の年÷12 と違う")
        plans = [p["計画額（円・税別）"] for p in seed["計画"] if p["期間"] == lab and p["状態"] == "有効"]
        if by_key["計画"]["月"][i] != (sum(plans) / 12 if plans else "未設定"):
            raise SystemExit(f"計画 の {m} が計画の行の年÷12 と違う")


def check_counts_integer_bps(action):
    """行動の目標件数を整数の bps で逆算し直し、build_yojitsu.py（浮動小数の ROUNDUP(ROUND(,6),0)）と同じかを確かめる（K-6）。"""
    top = {r["稼ぎ方"]: r for r in action["上の表"]}
    for row in action["目標件数と実績件数"]:
        t = top[row["稼ぎ方"]]
        if row["ステップ"] != "契約":
            continue
        plan = row["計画の額"]
        if not isinstance(plan, int):
            continue
        rate_bps = round(t["売上率"] * 10000)
        contract = -(-plan * 10000 // (t["1件あたりの額"] * rate_bps))
        proposal = -(-contract * 10000 // round(t["提案→契約の率"] * 10000))
        meeting = -(-proposal * 10000 // round(t["面談→提案の率"] * 10000))
        got = {r["ステップ"]: r["目標件数"] for r in action["目標件数と実績件数"]
               if r["稼ぎ方"] == row["稼ぎ方"] and r["期間"] == row["期間"]}
        if got != {"契約": contract, "提案": proposal, "面談": meeting}:
            raise SystemExit(f"行動の目標件数が整数の bps の逆算と違う: {row['稼ぎ方']} {row['期間']}")


# ---------------------------------------------------------------- (1') 予実の整数（#14 が照合する値）
INT_RULES = [
    "種から build_yojitsu.py を使わずに独立に計算した、予実の4段の整数の期待値。統合システムの純関数（core の money.md、要件 FR-REV-YOJITSU-004・007）と"
    "同じ決まりで計算し、#14（予実の API）はこの節と不一致0を確かめる。端数の決まりは要件の要確認11（代表が確かめる）",
    "月は期間（1〜5年目）の60か月（YYYY-MM）。予実の行の最後の月（期間の外）は持たない",
    "料率の額（着地予測・確度加重予測・実績）は、案件×月ごとに税別の額を足してから1回だけ 0 に向かって切り捨てる（確度加重は 額×売上率×確度 を1回で切り捨てる）。"
    "率は案件の売上率（制作受託は受託の契約の版の受託費率で、試験は種の案件の売上率で契約の版を作る）",
    "比（達成率・進み）は bp の整数（分子×10000÷分母 を 0 に向かって切り捨て）。分母が0以下か、分子・分母のどちらかが null なら null。"
    "行動の月の達成率は 実績件数×100×10000÷月の目標件数（100倍の整数）",
    "目標・計画の月の額は 年の額÷12 を整数で割り、端数は各年目の最後の月に寄せる。月の目標件数は 年の件数×100 を同じく割った100倍の整数",
    "値の元になる行が無い月・期間は null（0 で埋めない。未着・未到来）。流れの値（予測・入金・件数）の期間と通算は値のある月の和、"
    "年の値（目標・計画・目標件数）の通算は5つの期間がそろったときだけ和",
    "行動の目標件数は 契約＝計画÷（1件の額×既定の売上率）、提案＝契約÷提案→契約の率、面談＝提案÷面談→提案の率 をそれぞれ切り上げ（整数の bps）",
    "生成のときに Excel の値（予実.行・行動）と次のとおり合うことを確かめた: 料率の額は 0 ≤ 小数−整数 < 切り捨てた回数、目標・計画の月の額は差が12円未満、"
    "件数の年の値は同じ、月の目標件数は 小数×100 との差が12未満、整数が null のセルは Excel の 0・空・未設定",
]
YEN_TRUNC_KEYS = ("着地", "加重", "実績")  # 料率の額の行（稼ぎ方ごとの 着地/… 実績/… も）


def to_bps(x) -> int:
    """0.8 のような小数の率を、文字のまま bp の整数に読む（浮動小数を掛けない）。"""
    f = Fraction(str(x)) * 10000
    if f.denominator != 1:
        raise SystemExit(f"率 {x} は bp の整数で表せない")
    return int(f)


def tdiv(n: int, d: int) -> int:
    """0 に向かって切り捨てる割り算（統合システムの truncate と同じ）。"""
    q = abs(n) // abs(d)
    return q if (n >= 0) == (d > 0) else -q


def ratio_bps(n, d):
    return None if n is None or d is None or d <= 0 else tdiv(n * 10000, d)


def split12(total: int) -> list[int]:
    base = total // 12
    return [base] * 11 + [total - base * 11]


def ceil_div(n: int, d: int) -> int:
    return -(-n // d)


def add(a, b):
    return b if a is None else a + b


def sum_present(values):
    out = None
    for v in values:
        if v is not None:
            out = add(out, v)
    return out


def sum_all(values):
    return sum(values) if values and all(v is not None for v in values) else None


def integer_section(seed, yojitsu_rows, action):
    st = seed["設定"]
    vocab = seed["vocabulary"]
    periods = [(p["期間"], p["始め"][:7]) for p in st["期間"]]

    def months_of(start):
        y, m = map(int, start.split("-"))
        return [f"{y + (m - 1 + k) // 12:04d}-{(m - 1 + k) % 12 + 1:02d}" for k in range(12)]

    months = [m for _, start in periods for m in months_of(start)]
    mi = {m: i for i, m in enumerate(months)}
    n = len(months)
    slices = [(k * 12, (k + 1) * 12) for k in range(len(periods))]
    proj = {p["案件コード"]: p for p in seed["案件"]}
    stage_bps = {s["name"]: s["probability_bps"] for s in vocab["stages"]}
    earnings = [t["name"] for t in vocab["earning_types"]]
    default_rate = {t["name"]: t["default_sales_rate_bps"] for t in vocab["earning_types"]}
    sales_kubun = {k["name"] for k in vocab["kubun"] if k["direction"] == "入金" and k["sales_source"] == "はい"}
    thr = to_bps(st["着地確度"])
    scenario = st["採用シナリオ"]

    def flow(values, units=None):
        periods_ = [sum_present(values[a:b]) for a, b in slices]
        row = {"期間": periods_, "通算": sum_present(periods_), "月": values}
        if units is not None:
            row["_units"] = units
        return row

    def annual(per_period):
        """per_period: 期間ごとの (年の値, 12か月の値) か None。"""
        vals = [None if x is None else x[0] for x in per_period]
        ms = []
        for x in per_period:
            ms += [None] * 12 if x is None else x[1]
        return {"期間": vals, "通算": sum_all(vals), "月": ms}

    def ratio(num, den, scale=1):
        return {"期間": [ratio_bps(a, b) for a, b in zip(num["期間"], den["期間"])], "通算": ratio_bps(num["通算"], den["通算"]),
                "月": [ratio_bps(None if a is None else a * scale, b) for a, b in zip(num["月"], den["月"])]}

    # 予測: 採用シナリオか単一・予測に使う・売上の元の区分の行を、案件×月で足してから率を掛ける
    fsum = {}
    for f in seed["予測"]:
        if f["予測に使う"] != "はい" or f["シナリオ"] not in (scenario, "単一") or f["区分"] not in sales_kubun:
            continue
        i = mi.get(f["月"][:7])
        if i is not None:
            fsum[(f["案件コード"], i)] = fsum.get((f["案件コード"], i), 0) + f["金額（円・税別）"]
    landing, weighted = [None] * n, [None] * n
    landing_u, weighted_u = [0] * n, [0] * n
    landing_e = {e: [None] * n for e in earnings}
    landing_eu = {e: [0] * n for e in earnings}
    for (code, i), amount in sorted(fsum.items()):
        p = proj[code]
        rate, prob = p["rate_bps"], stage_bps[p["段階"]]
        weighted[i] = add(weighted[i], tdiv(amount * rate * prob, 10 ** 8))
        weighted_u[i] += 1
        if prob >= thr:
            s = tdiv(amount * rate, 10000)
            landing[i] = add(landing[i], s)
            landing_u[i] += 1
            landing_e[p["稼ぎ方"]][i] = add(landing_e[p["稼ぎ方"]][i], s)
            landing_eu[p["稼ぎ方"]][i] += 1

    # 実績: 入金の行（売上の元の区分）を案件×月で足してから率を掛ける。入金の総額は率を掛ける前
    asum = {}
    for r in seed["実績"]:
        if r["区分"] not in sales_kubun:
            continue
        if r["案件コード"] not in proj:
            raise SystemExit(f"案件に結びつかない入金の扱いを決めてから整数の期待値を作る: {r['実績ID']}")
        i = mi.get(r["日付"][:7])
        if i is not None:
            asum[(r["案件コード"], i)] = asum.get((r["案件コード"], i), 0) + r["金額（円・税別）"]
    actual, inflow = [None] * n, [None] * n
    actual_u = [0] * n
    actual_e = {e: [None] * n for e in earnings}
    actual_eu = {e: [0] * n for e in earnings}
    for (code, i), amount in sorted(asum.items()):
        p = proj[code]
        s = tdiv(amount * p["rate_bps"], 10000)
        inflow[i] = add(inflow[i], amount)
        actual[i] = add(actual[i], s)
        actual_u[i] += 1
        actual_e[p["稼ぎ方"]][i] = add(actual_e[p["稼ぎ方"]][i], s)
        actual_eu[p["稼ぎ方"]][i] += 1

    # 目標（有効な版・全体）と計画（有効な版）
    def target_cell(kind, label):
        rows = [t for t in seed["目標"] if t["目標の種類"] == kind and t["期間"] == label and t["稼ぎ方"] == "全体" and t["状態"] == "有効"]
        if len(rows) > 1:
            raise SystemExit(f"有効な目標が2つある: {kind} {label}")
        if not rows:
            return None
        amount = rows[0]["目標額（円）"]
        return (amount, split12(amount)) if rows[0]["単位"] == "期間の合計" else (amount * 12, [amount] * 12)

    def plan_amount(label, earning=None):
        rows = [p for p in seed["計画"] if p["期間"] == label and p["状態"] == "有効" and (earning is None or p["稼ぎ方"] == earning)]
        return sum(p["計画額（円・税別）"] for p in rows) if rows else None

    def plan_cell(label, earning=None):
        a = plan_amount(label, earning)
        return None if a is None else (a, split12(a))

    out = {}
    out["目標本人"] = {"単位": "円", **annual([target_cell("本人", lab) for lab, _ in periods])}
    out["目標会社"] = {"単位": "円", **annual([target_cell("会社", lab) for lab, _ in periods])}

    # 行動: 目標件数（稼ぎ方×ステップ×期間）と実績件数（ステップ×月、稼ぎ方×ステップ×期間）
    act = {a["稼ぎ方"]: a for a in seed["行動"]}
    counts = {}
    for e in earnings:
        a = act[e]
        unit, rate = a["1件あたりの額（円・税別）"], default_rate[e]
        to_contract, to_proposal = to_bps(a["提案→契約の率"]), to_bps(a["面談→提案の率"])
        for lab, _ in periods:
            plan = plan_amount(lab, e)
            if plan is None:
                counts[(e, lab)] = None
                continue
            contract = ceil_div(plan * 10000, unit * rate)
            proposal = ceil_div(contract * 10000, to_contract)
            counts[(e, lab)] = {"面談": ceil_div(proposal * 10000, to_proposal), "提案": proposal, "契約": contract}
    by_month, by_period = {}, {}
    for a in seed["活動記録"]:
        step = a["ステップ"]
        if step == "その他":
            continue
        e = a["稼ぎ方"] or proj[a["案件コード"]]["稼ぎ方"]
        ym = a["日付"][:7]
        lab = next((label for label, start in periods if ym in months_of(start)), None)
        if ym in mi:
            by_month[(step, mi[ym])] = by_month.get((step, mi[ym]), 0) + 1
        if lab:
            by_period[(e, step, lab)] = by_period.get((e, step, lab), 0) + 1
    action_rows = []
    for e in earnings:
        for step in ("面談", "提案", "契約"):
            for lab, _ in periods:
                c = counts[(e, lab)]
                target = None if c is None else c[step]
                got = by_period.get((e, step, lab))
                action_rows.append({"稼ぎ方": e, "ステップ": step, "期間": lab, "計画の額": plan_amount(lab, e), "目標件数": target, "実績件数": got,
                                    "達成率bp": ratio_bps(got, target)})
    for step in ("面談", "提案", "契約"):
        per = []
        for lab, _ in periods:
            cs = [counts[(e, lab)] for e in earnings]
            total = None if any(c is None for c in cs) else sum(c[step] for c in cs)
            per.append(None if total is None else (total, split12(total * 100)))
        target = {"単位": "件数（月は100倍）", **annual(per)}
        actual_counts = {"単位": "件数", **flow([by_month.get((step, i)) for i in range(n)])}
        out[f"行動目標/{step}"] = target
        out[f"行動実績/{step}"] = actual_counts
        out[f"行動率/{step}"] = {"単位": "bp", **ratio(actual_counts, target, 100)}

    plan_row = {"単位": "円", **annual([plan_cell(lab) for lab, _ in periods])}
    landing_row = {"単位": "円", **flow(landing, landing_u)}
    actual_row = {"単位": "円", **flow(actual, actual_u)}
    out["計画"] = plan_row
    out["率計画"] = {"単位": "bp", **ratio(plan_row, out["目標本人"])}
    out["着地"] = landing_row
    out["率着計"] = {"単位": "bp", **ratio(landing_row, plan_row)}
    out["加重"] = {"単位": "円", **flow(weighted, weighted_u)}
    out["実績"] = actual_row
    out["率実計"] = {"単位": "bp", **ratio(actual_row, plan_row)}
    out["率実績"] = {"単位": "bp", **ratio(actual_row, out["目標本人"])}
    out["入金"] = {"単位": "円", **flow(inflow)}
    out["不明入金"] = {"単位": "円", **flow([None] * n)}
    for e in earnings:
        pe = {"単位": "円", **annual([plan_cell(lab, e) for lab, _ in periods])}
        ae = {"単位": "円", **flow(actual_e[e], actual_eu[e])}
        out[f"計画/{e}"] = pe
        out[f"着地/{e}"] = {"単位": "円", **flow(landing_e[e], landing_eu[e])}
        out[f"実績/{e}"] = ae
        out[f"率実計/{e}"] = {"単位": "bp", **ratio(ae, pe)}

    top = [{"稼ぎ方": e, "1件あたりの額": act[e]["1件あたりの額（円・税別）"], "売上率bp": default_rate[e],
            "1件の売上": tdiv(act[e]["1件あたりの額（円・税別）"] * default_rate[e], 10000),
            "面談→提案bp": to_bps(act[e]["面談→提案の率"]), "提案→契約bp": to_bps(act[e]["提案→契約の率"])} for e in earnings]

    check_integers_against_excel(out, yojitsu_rows, action, top, action_rows, months, [lab for lab, _ in periods])
    keys = [r["鍵"] for r in yojitsu_rows]
    if sorted(keys) != sorted(out):
        raise SystemExit(f"整数の行の鍵が予実の行と違う: {sorted(set(keys) ^ set(out))}")
    # 目標・計画から出る行（予実の行で fictional: true の行と、計画の額・目標件数を持つ行動の表の行）に fictional: true を付ける
    # （予実の行・行動の表と同じ印。行動指標の計画 段B #23 の test/fictional-compensation.test.mjs が、架空の目標・計画の額は印のある行にだけ出ることを見る）
    fictional_keys = {r["鍵"] for r in yojitsu_rows if r.get("fictional") is True}
    rows = []
    for key in keys:
        row = {k: v for k, v in out[key].items() if not k.startswith("_")}
        rows.append({"鍵": key, **row, **({"fictional": True} if key in fictional_keys else {})})
    action_rows = [{**r, "fictional": True} for r in action_rows]
    return {"決まり": INT_RULES, "月": months, "期間": [lab for lab, _ in periods], "行": rows, "行動": action_rows, "行動の上の表": top}


def check_integers_against_excel(out, yojitsu_rows, action, top, action_rows, months, labels):
    """整数の期待値が、build_yojitsu.py の小数の値（Excel と同じ規則）と、端数の決まりの範囲で合うかを確かめる。"""
    flt = {r["鍵"]: r for r in yojitsu_rows}
    zero_like = (None, 0, 0.0, "", "未設定")

    def fail(msg):
        raise SystemExit(f"整数の期待値が Excel の値と合わない: {msg}")

    for key, row in out.items():
        f = flt[key]
        if len(f["月"]) != len(months) + 1 or f["月"][-1] not in zero_like:
            fail(f"{key} の期間の外の月")
        cells = [(f"{months[i]}", row["月"][i], f["月"][i], (row.get("_units") or [None] * len(months))[i]) for i in range(len(months))]
        units = row.get("_units")
        for k, lab in enumerate(labels):
            u = None if units is None else sum(units[k * 12:(k + 1) * 12])
            cells.append((lab, row["期間"][k], f["期間"][lab], u))
        cells.append((TOTAL, row["通算"], f["期間"][TOTAL], None if units is None else sum(units)))
        for where, iv, fv, u in cells:
            if iv is None:
                if fv not in zero_like:
                    fail(f"{key} {where}: 整数は null で Excel は {fv}")
                continue
            if not isinstance(fv, (int, float)):
                fail(f"{key} {where}: 整数は {iv} で Excel は {fv}")
            unit = row["単位"]
            if unit == "bp":
                continue  # 比は整数の行どうしから作る（上の ratio）。分母の Excel の空は整数でも null（上の iv is None の枝）
            if unit == "件数":
                if fv != iv:
                    fail(f"{key} {where}: 件数 {iv} と {fv}")
            elif unit.startswith("件数（"):
                if where in labels or where == TOTAL:
                    if fv != iv:
                        fail(f"{key} {where}: 年の件数 {iv} と {fv}")
                elif abs(Fraction(str(fv)) * 100 - iv) >= 12:
                    fail(f"{key} {where}: 月の件数（100倍）{iv} と {fv}")
            elif key.split("/")[0] in YEN_TRUNC_KEYS:
                d = Fraction(str(fv)) - iv
                if not (0 <= d < max(u or 0, 1)):
                    fail(f"{key} {where}: 切り捨ての差 {float(d)}（切り捨て {u} 回）")
            elif key == "入金":
                if fv != iv:
                    fail(f"{key} {where}: {iv} と {fv}")
            else:  # 目標・計画（年の値は同じ、月は年÷12 との差が12円未満）
                if (where in labels or where == TOTAL) and fv != iv:
                    fail(f"{key} {where}: 年の額 {iv} と {fv}")
                if abs(Fraction(str(fv)) - iv) >= 12:
                    fail(f"{key} {where}: 月の額 {iv} と {fv}")
    for t in top:
        g = next(x for x in action["上の表"] if x["稼ぎ方"] == t["稼ぎ方"])
        if g["1件あたりの額"] != t["1件あたりの額"] or to_bps(g["売上率"]) != t["売上率bp"] or not (0 <= Fraction(str(g["1件の売上"])) - t["1件の売上"] < 1):
            fail(f"行動の上の表 {t['稼ぎ方']}")
    for r in action_rows:
        g = next(x for x in action["目標件数と実績件数"] if (x["稼ぎ方"], x["ステップ"], x["期間"]) == (r["稼ぎ方"], r["ステップ"], r["期間"]))
        if g["計画の額"] != r["計画の額"] or g["目標件数"] != r["目標件数"] or (g["実績件数"] or None) != r["実績件数"]:
            fail(f"行動 {r['稼ぎ方']} {r['ステップ']} {r['期間']}")


# ---------------------------------------------------------------- (3) 利益額（K-16）
def profit_details(seed, invoices_doc):
    """基準×月×案件×費目の明細。売上は税別の額を足してから売上率を掛けて切り捨てる。"""
    proj = {p["案件コード"]: p for p in seed["案件"]}
    kubun = {k["name"]: k for k in seed["vocabulary"]["kubun"]}
    with_work = {w["project_code"] for w in seed["works"]}
    work_project = {w["work_id"]: w["project_code"] for w in seed["works"]}
    withholding = {c["partner_code"] for c in invoices_doc["contracts"] if c["foreign_withholding_bps"]}
    invoice_partner = {v["invoice_number"]: v["partner_code"] for v in invoices_doc["invoices"]}
    as_of = seed["as_of"]
    raw = {}   # (基準, 月, 案件コード, 費目) → 税別の額の和
    unknown = {}  # (基準, 月) → 案件に結べない入金の税別の額

    def add(basis, month, code, item, amount):
        if item not in ITEMS:
            raise SystemExit(f"費目が決まっていない: {item}")
        raw[(basis, month, code, item)] = raw.get((basis, month, code, item), 0) + amount

    def sale_item(code, partner):
        if partner in withholding:
            return "売上（源泉のある取引先）"
        if code not in with_work:
            return "売上（作品の無い案件）"
        return "売上"

    # 入金の基準の売上と、実績の出金
    pay_index = {}
    for x in seed["investments"]:
        code = next(w["project_code"] for w in seed["works"] if w["work_code"] == x["work_code"])
        for p in x["payments"]:
            pay_index.setdefault((code, p["paid_on"], p["amount_yen"]), x["kind"])
    for r in seed["実績"]:
        k = kubun.get(r["区分"])
        if k is None:
            raise SystemExit(f"実績の区分が無い: {r['実績ID']}")
        code, month, amount = r["案件コード"], ym(r["日付"]), r["金額（円・税別）"]
        if k["direction"] == "入金":
            if k["sales_source"] != "はい":
                raise SystemExit(f"売上に数えない入金（{r['区分']}）の扱いは U-5 の残り。決めてから golden を作る: {r['実績ID']}")
            if not code or code not in proj:
                unknown[("cash", month)] = unknown.get(("cash", month), 0) + amount
                continue
            add("cash", month, code, sale_item(code, invoice_partner.get(r.get("invoice_number"))), amount)
        elif r["区分"] == "出資":
            if (code, r["日付"], amount) not in pay_index:
                raise SystemExit(f"出資の実績の行が出資の記録の払込に無い: {r['実績ID']}")
            for basis in ("cash", "accrual"):
                add(basis, month, code, "出資の実績の行（出資の記録の払込と同じ）", amount)
        elif r["区分"] == "買付":
            item = "買付の MG" if "MG" in r["項目"] else "買い切りの買付額"
            for basis in ("cash", "accrual"):
                add(basis, month, code, item, amount)
        elif r["区分"] == "制作の実費":
            for basis in ("cash", "accrual"):
                add(basis, month, code, "制作の実費", amount)
        else:
            raise SystemExit(f"扱いを決めていない出金の区分（{r['区分']}）。決めてから golden を作る: {r['実績ID']}")

    # 計上の基準の売上（請求明細の計上月）
    for v in invoices_doc["invoices"]:
        for line in v["lines"]:
            code = work_project.get(line["work_id"])
            month = line["accounting_month"]
            if code is None or code not in proj:
                unknown[("accrual", month)] = unknown.get(("accrual", month), 0) + line["amount_ex_tax"]
                continue
            add("accrual", month, code, sale_item(code, v["partner_code"]), line["amount_ex_tax"])

    # 出資の記録の払込（J-11b (a) の4費目のうち2つ）。種の時点より後の払込は、まだ払っていないので入れない
    for x in seed["investments"]:
        code = next(w["project_code"] for w in seed["works"] if w["work_code"] == x["work_code"])
        item = "自社出資の出資額" if x["kind"] == "自社出資" else "製作委員会の出資額（自社の分）"
        if x["kind"] not in ("自社出資", "製作委員会") or x["tax_basis"] != "税抜":
            raise SystemExit(f"出資の記録の種類か税の基準が想定と違う: {x['investment_id']}")
        for p in x["payments"]:
            if p["paid_on"] <= as_of:
                for basis in ("cash", "accrual"):
                    add(basis, ym(p["paid_on"]), code, item, p["amount_yen"])

    # 経費（制作費の実費と P&A）。計上は計上月、入金は払った日（未払は入れない）
    for x in seed["expenses"]:
        if x["category"] == "制作費":
            item = "費用区分「制作費」の経費"
        elif x["category"] in ("宣伝費", "配給費"):
            item = f"{x['category']}（負担者 {'自社' if x['bearer'] == '自社' else '自社以外'}）"
        else:
            raise SystemExit(f"扱いを決めていない費用区分: {x['expense_id']}")
        add("accrual", ym(x["accounting_month"]), x["project_code"], item, x["amount_ex_tax"])
        if x["paid_on"] and x["paid_on"] <= as_of:
            add("cash", ym(x["paid_on"]), x["project_code"], item, x["amount_ex_tax"])

    details = []
    for (basis, month, code, item), amount in sorted(raw.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2], ITEM_ORDER.index(kv[0][3]))):
        side, status, _ = ITEMS[item]
        row = {"基準": basis, "月": month, "案件コード": code, "稼ぎ方": proj[code]["稼ぎ方"], "費目": item, "向き": side,
               "状態": status, "税別の額": amount}
        if side == "売上":
            row["売上率bps"] = proj[code]["rate_bps"]
            row["額"] = rate_amount(amount, proj[code]["rate_bps"])
        else:
            row["額"] = amount
        details.append(row)
    return details, unknown


def profit_section(seed, invoices_doc):
    details, unknown = profit_details(seed, invoices_doc)
    months = sorted({d["月"] for d in details} | {m for _, m in unknown})
    idx = {m: i for i, m in enumerate(months)}
    bases = {}
    for basis in ("cash", "accrual"):
        rows = [d for d in details if d["基準"] == basis]
        sales, cost = [0] * len(months), [0] * len(months)
        for d in rows:
            if d["状態"] != "確定":
                continue
            (sales if d["向き"] == "売上" else cost)[idx[d["月"]]] += d["額"]
        profit = [s - c for s, c in zip(sales, cost)]
        groups = {}
        for d in rows:
            if d["状態"] == "確定":
                continue
            g = groups.setdefault(d["費目"], {"費目": d["費目"], "状態": d["状態"], "向き": d["向き"], "理由": ITEMS[d["費目"]][2],
                                               "額": [0] * len(months), "合計": 0})
            g["額"][idx[d["月"]]] += d["額"]
            g["合計"] += d["額"]
        by_project = {}
        for d in rows:
            if d["状態"] != "確定":
                continue
            p = by_project.setdefault(d["案件コード"], {"案件コード": d["案件コード"], "稼ぎ方": d["稼ぎ方"], "売上": 0, "経費": 0})
            p["売上" if d["向き"] == "売上" else "経費"] += d["額"]
        for p in by_project.values():
            p["利益額"] = p["売上"] - p["経費"]
        unknown_row = [unknown.get((basis, m), 0) for m in months]
        bases[basis] = {
            "確定": {"売上": sales, "経費": cost, "利益額": profit,
                     "合計": {"売上": sum(sales), "経費": sum(cost), "利益額": sum(profit)}},
            "未確定": [groups[k] for k in ITEM_ORDER if k in groups and groups[k]["状態"] == "未確定"],
            "除外": [groups[k] for k in ITEM_ORDER if k in groups and groups[k]["状態"] == "除外"],
            "案件不明": {"理由": "案件に結べない入金（計上は請求明細）。売上に数えず、別の行に出す", "額": unknown_row,
                         "合計": sum(unknown_row)},
            "案件別": [by_project[k] for k in sorted(by_project)],
        }
    return {
        "式": "利益額＝売上（入金×売上率。計上の基準は計上の売上×売上率）− 案件の経費（K-16）",
        "決まり": RULES,
        "費目": [{"費目": k, "向き": v[0], "状態": v[1], "理由": v[2]} for k, v in ITEMS.items()],
        "月": months,
        "基準": bases,
        "明細": details,
    }


# ---------------------------------------------------------------- 値の無い経路
def empty_paths(profit, yojitsu, earn):
    """種に行が無く golden では値を確かめられない経路を、(1)(3) の値から数える（手で書かない）。"""
    yrow = {r["鍵"]: r for r in yojitsu}
    bases = ("cash", "accrual")
    return {
        "理由": EMPTY_REASON,
        "利益額の明細が無い費目": {b: [k for k in ITEM_ORDER if not any(d["基準"] == b and d["費目"] == k for d in profit["明細"])]
                              for b in bases},
        "利益額の案件不明が全月0の基準": [b for b in bases if all(v == 0 for v in profit["基準"][b]["案件不明"]["額"])],
        "予実の実績が全月0の稼ぎ方": [e for e in earn if all(not v for v in yrow[f"実績/{e}"]["月"])],
    }


# ---------------------------------------------------------------- ORG の値が入っていないこと
def fab_leaks(bm, golden_text: str, golden) -> list[str]:
    """build_yojitsu.py が ORG の種として持つ案件コード・名前・金額が golden に無いかを見る。値そのものは出さない。"""
    codes, names, amounts = set(), set(), set()
    for row in getattr(bm, "PROJECT_ROWS", []):
        codes.add(row[0])
        names.update(x for x in (row[1], row[3], row[10]) if isinstance(x, str) and len(x) >= 4)
        if isinstance(row[6], int):
            amounts.add(row[6])
    for row in getattr(bm, "TARGET_ROWS", []) + getattr(bm, "PLAN_ROWS", []):
        codes.add(row[0])
        amounts.update(v for v in row if isinstance(v, int) and not isinstance(v, bool) and v >= 100_000)
    for v in getattr(bm, "ACTION_PER", {}).values():
        amounts.add(v[0])
    for row in getattr(bm, "ACTIVITY_ROWS", []):
        names.update(x for x in row if isinstance(x, str) and len(x) >= 8)
    names.update(getattr(bm, "TRIAL_PROJECT", {}).keys())
    found = []
    if any(re.search(rf"(?<![A-Za-z0-9-]){re.escape(c)}(?![0-9])", golden_text) for c in codes):
        found.append("ORG の種の案件・目標・計画のコード")
    if any(n in golden_text for n in names):
        found.append("ORG の種の名前・相手・企画フォルダ・活動の文")

    def walk(v):
        if isinstance(v, bool):
            return
        if isinstance(v, (int, float)) and float(v).is_integer() and int(v) in amounts:
            found.append("ORG の種の金額")
        elif isinstance(v, dict):
            for x in v.values():
                walk(x)
        elif isinstance(v, list):
            for x in v:
                walk(x)

    walk(golden)
    return sorted(set(found))


# ---------------------------------------------------------------- 書き出し
def is_rowlike(v) -> bool:
    def scalar(x):
        return not isinstance(x, (dict, list))
    if isinstance(v, dict):
        return all(scalar(x) or (isinstance(x, list) and all(scalar(y) for y in x))
                   or (isinstance(x, dict) and all(scalar(y) for y in x.values())) for x in v.values())
    if isinstance(v, list):
        return all(scalar(x) for x in v)
    return True


def render(v, depth=0) -> str:
    """上の2段と、行より大きい入れ物を開き、行（値が数・文字・数の並びだけのもの）は1行に1件で書く。"""
    pad = "  " * depth
    if isinstance(v, dict) and v and (depth < 2 or not is_rowlike(v)):
        body = ",\n".join(f"{pad}  {json.dumps(k, ensure_ascii=False)}: {render(x, depth + 1)}" for k, x in v.items())
        return "{\n" + body + "\n" + pad + "}"
    sentences = isinstance(v, list) and any(isinstance(x, str) and len(x) > 20 for x in v)
    if isinstance(v, list) and v and (sentences or not is_rowlike(v)):
        body = ",\n".join(f"{pad}  {render(x, depth + 1)}" for x in v)
        return "[\n" + body + "\n" + pad + "]"
    return json.dumps(v, ensure_ascii=False)


def make(build_path: Path, keep_build: Path | None):
    seed_path, inv_path = HERE / "seed.json", RM_DIR / "invoices.json"
    seed = json.loads(lf_bytes(seed_path).decode("utf-8"))
    invoices_doc = json.loads(lf_bytes(inv_path).decode("utf-8"))
    if seed.get("fictional") is not True or invoices_doc.get("fictional") is not True:
        raise SystemExit("種か入金の確認の請求に fictional: true が無い")
    bm = load_build_yojitsu(build_path)
    guard = protected_files(bm)
    before = stat_of(guard)
    earn_note = use_seed_earnings(bm, seed["vocabulary"])
    work = Path(tempfile.mkdtemp(prefix="yojitsu-golden-"))
    try:
        produced, data, wb, g_main, g_month, stage_layout, out = run_build(bm, seed, work)
        if keep_build:
            keep_build.mkdir(parents=True, exist_ok=True)
            for name in ("予実管理表.xlsx", "expected.json"):
                shutil.copyfile(out / name, keep_build / name)
    finally:
        shutil.rmtree(work, ignore_errors=True)
    if stat_of(guard) != before:
        raise SystemExit("ORG の予実管理表.xlsx・expected.json・trials_snapshot.json の大きさか更新時刻が変わった")

    cells = produced["sheets"]
    main_derived = target_derived(g_main.specs, {("行動目標", s) for s in bm.STEPS})
    month_derived = target_derived(g_month.specs, set())
    yojitsu = grid_rows(bm, wb["予実"], g_main, cells["予実"], main_derived)
    monthly = grid_rows(bm, wb["月別"], g_month, cells["月別"], month_derived)
    action = action_section(bm, data, cells["行動"])
    check_targets_against_build(bm, seed, yojitsu)
    check_counts_integer_bps(action)
    integers = integer_section(seed, yojitsu, action)

    profit = profit_section(seed, invoices_doc)
    golden = {
        "fictional": True,
        "notice": NOTICE,
        "generator": "golden.py",
        "plan": PLAN,
        "sources": {"yojitsu/seed.json": sha256(seed_path), "receipt-matching/invoices.json": sha256(inv_path)},
        "build_yojitsu": {"ファイル": "会社ルート ORG/事業計画/_build/build_yojitsu.py（#23b）",
                          "走らせ方": "--seed（seed.json の入力の表を列の名前だけに絞った写し）と --out（一時フォルダ）。--snapshot なし",
                          **earn_note},
        "notes": [
            "(1) 予実・行動・段階・月別は build_yojitsu.py が出した expected.json（Excel の式と同じ規則の値）を、行の鍵・期間・月で引ける形に"
            "直したもの。額は Excel と同じく丸めない浮動小数。null は Excel の空（期間の外・分母が0）、「未設定」は目標か計画が無い月",
            "(1) の額の端数: 実績・着地・確度加重・目標と計画の月の額（年÷12）などは丸めない小数（Excel と同じ）。整数円で計算する"
            "統合システムの純関数（core の money.md）と #14 の予実の API が照合するのは (1') 予実の整数（種から独立に計算した整数円・bp・100倍の件数の期待値。"
            "決まりは同じ節。端数の決まりは要件の要確認11 で代表が確かめる）。入金の基準の売上を案件×月で切り捨てた整数は 利益額.明細 にもある",

            "予実の4段は入金に固定（D-039）。実績は入金の行×案件の売上率、着地は確度が設定の着地確度以上の案件の採用シナリオの売上額",
            "fictional: true の行は目標・計画（種の fictional: true の架空額）から出た値",
            "(2) 施策の送信後の期待値は #20 で施策の入力値を固定したあとに作る（送信後 の枠）",
            "(3) 利益額は build_yojitsu.py を使わず、種と入金の確認の請求から計算した。決まりは 利益額.決まり",
            "値の無い経路 は、種に行が無く golden では確かめられない経路（#23e で4区分の行を足したあとに作り直す）",
        ],
        "稼ぎ方": list(bm.EARN),
        "期間": seed["設定"]["期間"],
        "月": [m.isoformat()[:7] for m in bm.MONTHS],
        "設定": {k: seed["設定"][k] for k in ("採用シナリオ", "締めた月", "着地確度", "案件数目標", "月別の予測範囲")},
        "目標・計画": targets_section(seed),
        "予実": {"行": yojitsu},
        "予実の整数": integers,
        "行動": action,
        "段階": stage_section(bm, wb["段階"], cells["段階"], stage_layout),
        "月別": {"行": monthly},
        "印": flags_section(produced, data),
        "利益額": profit,
        "値の無い経路": empty_paths(profit, yojitsu, list(bm.EARN)),
        "送信後": {"状態": "未作成", "依存": "#20",
                   "理由": "施策の入力値（K-17 の7項目）は #20 の計測のレシピで固定する。固定したあとに、種＋その施策から独立に計算して埋める",
                   "施策": None,
                   "値": {"予実の4段": None, "タスク": None, "試算": None, "利益額": None, "説明の帳票": None}},
    }
    text = render(golden) + "\n"
    leaks = fab_leaks(bm, text, golden)
    if leaks:
        raise SystemExit(f"golden に ORG の値がある（書かない）: {leaks}")
    return text.encode("utf-8")


def main():
    ap = argparse.ArgumentParser(description="案件収支の予実の golden（架空版。expected.json）を作る")
    ap.add_argument("--build-yojitsu", help="会社ルートの ORG/事業計画/_build/build_yojitsu.py")
    ap.add_argument("--out", type=Path, default=HERE, help="expected.json を書くフォルダ（既定はこのフォルダ）")
    ap.add_argument("--keep-build", type=Path, help="build_yojitsu.py の出力（予実管理表.xlsx と expected.json）を残すフォルダ")
    ap.add_argument("--check", action="store_true", help="作り直して、置いてある expected.json と同じかを確かめる（書かない）")
    args = ap.parse_args()
    data = make(find_build_yojitsu(args.build_yojitsu), args.keep_build)
    if args.check:
        same = lf_bytes(HERE / "expected.json") == data
        print("expected.json は作り直したものと同じ" if same else "expected.json が作り直したものと違う（作り直してコミットする）")
        sys.exit(0 if same else 1)
    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / "expected.json").write_bytes(data)
    print(f"expected.json={hashlib.sha256(data).hexdigest()[:12]} bytes={len(data)}")


if __name__ == "__main__":
    main()
