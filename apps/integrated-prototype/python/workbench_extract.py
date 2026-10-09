"""Read CSV/XLSX values locally without recalculating formulas or external links."""
import base64, csv, datetime, hashlib, io, json, os, sys, zipfile

def value(v):
    if v is None: return ''
    if isinstance(v,datetime.datetime): return v.date().isoformat() if v.time()==datetime.time() else v.isoformat()
    if isinstance(v,datetime.date): return v.isoformat()
    if isinstance(v,float) and v.is_integer(): return int(v)
    return v

def extract(request):
    name=os.path.basename(str(request.get('name','')));extension=os.path.splitext(name)[1].lower()
    data=base64.b64decode(request['base64'],validate=True)
    if not data or len(data)>5*1024*1024: raise ValueError('ファイルは5MB以内にしてください')
    result={'name':name,'sourceHash':hashlib.sha256(data).hexdigest(),'byteLength':len(data),'extractorVersion':'workbench-extract-1'}
    if extension=='.csv':
        try: text=data.decode('utf-8-sig');encoding='utf-8-sig'
        except UnicodeDecodeError: text=data.decode('cp932');encoding='cp932'
        rows=[]
        for row in csv.reader(io.StringIO(text),strict=True):
            if len(rows)>=10001 or len(row)>100: raise ValueError('10001行・100列以内にしてください')
            rows.append(row)
        result.update(sourceEncoding=encoding,sheets=[{'name':'CSV','rows':rows,'formulaIssues':[]}])
    elif extension=='.xlsx':
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist())>40*1024*1024: raise ValueError('展開後のExcelが大きすぎます')
        values=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
        formulas=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        sheets=[]
        try:
            if len(values.worksheets)>30: raise ValueError('シート数は30以内にしてください')
            for sheet in values.worksheets:
                if (sheet.max_row or 0)>10001 or (sheet.max_column or 0)>100: raise ValueError('1シートは見出し込み10001行・100列以内にしてください')
                rows=[]
                for row in sheet.iter_rows(values_only=True):
                    if len(rows)>=10001 or len(row)>100: raise ValueError('10001行・100列以内にしてください')
                    rows.append([value(cell) for cell in row])
                issues=[]
                for row in formulas[sheet.title].iter_rows():
                    for cell in row:
                        if cell.data_type=='f' and (cell.row>len(rows) or cell.column>len(rows[cell.row-1]) or rows[cell.row-1][cell.column-1]==''):
                            issues.append({'cell':cell.coordinate,'reason':'数式の保存済み計算値がありません。Excelで再計算・保存してください。'})
                sheets.append({'name':sheet.title,'rows':rows,'formulaIssues':issues})
        finally: values.close();formulas.close()
        result['sheets']=sheets
    else: raise ValueError('CSVまたはXLSXを選んでください')
    for sheet in result['sheets']:
        if len(sheet['rows'])>10001 or any(len(row)>100 for row in sheet['rows']): raise ValueError('10001行・100列以内にしてください')
    return result

if __name__=='__main__':
    try: json.dump(extract(json.load(sys.stdin)),sys.stdout,ensure_ascii=False)
    except Exception as error: json.dump({'error':str(error)},sys.stdout,ensure_ascii=False);sys.exit(1)
