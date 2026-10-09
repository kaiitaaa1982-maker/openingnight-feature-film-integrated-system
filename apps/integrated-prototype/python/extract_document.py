import base64, contextlib, datetime, csv, hashlib, io, json, math, os, re, statistics, sys, unicodedata, zipfile

VERSION = "workflow-extractor-3"

def cell_text(value):
    if value is None: return ""
    if isinstance(value, datetime.datetime) and value.time() == datetime.time(0): return value.date().isoformat()
    if isinstance(value, datetime.date): return value.isoformat()
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value)

def rows_from_xlsx(data):
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return [{"name": ws.title, "rows": [[cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]} for ws in book.worksheets]

def text_from_docx(data):
    from docx import Document
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)

def text_from_pdf(data):
    return pdf_script(data)["text"]

def tables_from_pdf(data):
    try:
        import pdfplumber
        sheets = []
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for page_no, page in enumerate(pdf.pages, 1):
                    for table_no, table in enumerate(page.extract_tables() or [], 1):
                        rows = [["" if cell is None else str(cell) for cell in row] for row in table]
                        if rows: sheets.append({"name":f"PDF p{page_no} table {table_no}", "rows":rows})
        return sheets
    except Exception:
        return []

# ─── PDF の行（縦書き台本の組み直し） ─────────────────────────────
# 縦書きのPDFは1文字ずつ・または縦書きフォント（-V）で置かれ、そのまま読むと1文字1行になる。
# 文字の位置から列（縦の1行）を組み直し、右の列から順に並べる。余白の柱・ページ番号は落とす。
# 横書きのページは従来どおり pypdf の抽出文字列を使う。

JP = "\u3005\u3006\u3040-\u30ff\u3400-\u9fff\uf900-\ufaff"
JP_CHAR = re.compile(f"[{JP}]")

def _mul(a, b):
    return [a[0]*b[0]+a[1]*b[2], a[0]*b[1]+a[1]*b[3], a[2]*b[0]+a[3]*b[2], a[2]*b[1]+a[3]*b[3], a[4]*b[0]+a[5]*b[2]+b[4], a[4]*b[1]+a[5]*b[3]+b[5]]

def _fold(text):
    # 縦書き用の約物（﹁︒ 等）と康熙部首（⼭ 等）だけを通常の文字へ寄せる。ほかの文字は変えない
    return "".join(unicodedata.normalize("NFKC", ch) if "\ufe10" <= ch <= "\ufe4f" or "\u2e80" <= ch <= "\u2fdf" else ch for ch in text)

def _page_runs(page):
    runs = []
    def visit(text, cm, tm, font_dict, font_size):
        t = (text or "").replace("\r", "").replace("\n", "")
        if not t.strip(): return
        m = _mul(tm, cm)
        size = abs(font_size or 0) * (math.hypot(m[2], m[3]) or 1)
        encoding = ""
        try: encoding = str(font_dict.get("/Encoding", "")) if font_dict is not None else ""
        except Exception: encoding = ""
        runs.append({"x": m[4], "y": m[5], "text": _fold(t), "size": size, "vertical": encoding.endswith("-V")})
    page.extract_text(visitor_text=visit)
    return runs

def _near(points, size):
    # 近くの点だけを見るための格子（1文字ずつ置いたページでも全組み合わせを比べない）
    grid = {}
    for p in points: grid.setdefault((int(p["x"] // size), int(p["y"] // size)), []).append(p)
    def near(p):
        gx, gy = int(p["x"] // size), int(p["y"] // size)
        for dx in (-2, -1, 0, 1, 2):
            for dy in (-2, -1, 0, 1, 2):
                for q in grid.get((gx + dx, gy + dy), ()):
                    if q is not p: yield q
    return near

def _stacked(p, q, size):
    return abs(q["x"] - p["x"]) <= 0.45 * size and 0.6 * size <= abs(q["y"] - p["y"]) <= 1.4 * size

def _is_vertical_page(runs, size):
    vertical = sum(len(r["text"]) for r in runs if r["vertical"])
    horizontal = sum(len(r["text"]) for r in runs if not r["vertical"] and len(r["text"].strip()) > 1 and JP_CHAR.search(r["text"]))
    singles = [r for r in runs if not r["vertical"] and len(r["text"].strip()) == 1]
    near = _near(singles, size)
    for s in singles:
        others = list(near(s))
        if any(_stacked(s, t, size) for t in others): vertical += 1
        elif any(abs(t["y"] - s["y"]) <= 0.45 * size and 0.6 * size <= abs(t["x"] - s["x"]) <= 1.4 * size for t in others): horizontal += 1
    return vertical >= 5 and vertical > horizontal

def _vertical_columns(runs, size):
    tol = 0.45 * size
    for r in runs:
        r["kind"] = "v" if r["vertical"] else None
    singles = [r for r in runs if r["kind"] is None and len(r["text"].strip()) == 1]
    near = _near(singles, size)
    for s in singles:
        if any(_stacked(s, t, size) for t in near(s)): s["kind"] = "v"
    body = [r for r in runs if r["kind"] == "v"]
    if not body: return []
    extent = lambda r: r["y"] - (len(r["text"]) - 1) * r["size"] if r["vertical"] else r["y"]
    top, bottom = max(r["y"] for r in body), min(extent(r) for r in body)
    columns = []
    for r in sorted(body, key=lambda r: -r["x"]):
        if columns and abs(columns[-1]["x"] - r["x"]) <= tol:
            columns[-1]["runs"].append(r)
        else:
            columns.append({"x": r["x"], "runs": [r]})
    for c in columns:
        c["top"], c["bottom"] = max(r["y"] for r in c["runs"]), min(extent(r) for r in c["runs"])
    for r in runs:
        if r["kind"] is not None: continue
        near = [c for c in columns if abs(c["x"] - r["x"]) <= 0.6 * size]
        column = min(near, key=lambda c: abs(c["x"] - r["x"])) if near else None
        if column and column["bottom"] - 0.5 * size <= r["y"] <= column["top"] + 3 * size:
            column["runs"].append(r)          # 柱の上のシーン番号・縦中横の数字
        elif bottom - 0.5 * size <= r["y"] <= top + 0.5 * size and JP_CHAR.search(r["text"]):
            columns.append({"x": r["x"], "runs": [r], "top": r["y"], "bottom": r["y"]})  # 1文字だけの行
        # それ以外は余白（柱・ページ番号）として落とす
    columns.sort(key=lambda c: -c["x"])
    for c in columns:
        c["text"] = "".join(r["text"] for r in sorted(c["runs"], key=lambda r: -r["y"])).strip()
    return [c for c in columns if c["text"]]

def pdf_script(data):
    """PDFの本文を行の並びにする。戻り値: text（本文）、lines（[{text,page,slot}]）、layout（向き・1頁の行数）"""
    import logging
    from pypdf import PdfReader
    logging.getLogger("pypdf").setLevel(logging.ERROR)
    reader = PdfReader(io.BytesIO(data))
    pages, vertical_pages = [], 0
    for page_no, page in enumerate(reader.pages):
        runs = _page_runs(page)
        sizes = [r["size"] for r in runs if r["size"] > 0]
        size = statistics.median(sizes) if sizes else 10.0
        if runs and _is_vertical_page(runs, size):
            columns = _vertical_columns(runs, size)
            vertical_pages += 1
            pages.append({"vertical": True, "size": size, "columns": columns})
        else:
            text = page.extract_text() or ""
            pages.append({"vertical": False, "lines": [_fold(line) for line in text.replace("\r", "").split("\n")]})
    lines, texts = [], []
    xs_right, xs_left, pitches = [], [], []
    for p in pages:
        if p["vertical"] and len(p["columns"]) >= 5:
            xs = [c["x"] for c in p["columns"]]
            xs_right.append(max(xs)); xs_left.append(min(xs))
            pitches += [a - b for a, b in zip(xs, xs[1:]) if a - b >= 0.8 * p["size"]]
    pitch = statistics.median(pitches) if pitches else None
    right = statistics.median(xs_right) if xs_right else None
    slots_vertical = round((statistics.median(xs_right) - statistics.median(xs_left)) / pitch) + 1 if pitch and xs_left else None
    horizontal_counts = []
    for page_no, p in enumerate(pages):
        if p["vertical"]:
            for c in p["columns"]:
                slot = max(0, round((right - c["x"]) / pitch)) if pitch and right is not None else None
                lines.append({"text": c["text"], "page": page_no, "slot": slot})
                texts.append(c["text"])
        else:
            kept = p["lines"]
            horizontal_counts.append(len(kept))
            for index, line in enumerate(kept):
                lines.append({"text": line, "page": page_no, "slot": index})
                texts.append(line)
        texts.append("")
    if vertical_pages and vertical_pages >= len(pages) / 2:
        layout = {"orientation": "vertical", "slotsPerPage": slots_vertical}
    else:
        layout = {"orientation": "horizontal", "slotsPerPage": max(horizontal_counts) if horizontal_counts else None}
    layout["pages"] = len(pages)
    return {"text": "\n".join(texts).strip("\n"), "lines": lines, "layout": layout}

# ─── 台本の読み取り（シーン見出し・昼夜・出演・頁） ──────────────────────
TIME_WORDS = {"N": ("夜", "深夜", "NIGHT", "N"), "D": ("昼", "日中", "朝", "早朝", "翌朝", "DAY", "D"), "DN": ("夕", "夕方", "夕刻", "薄暮", "夜明け", "明け方")}
TIME_TOKEN = "深夜|夜明け|明け方|日中|早朝|翌朝|夕方|夕刻|薄暮|夜|昼|朝|夕|NIGHT|DAY|D|N"
GLYPH_HEADING = re.compile(r"^\s*[○◯〇](?![○◯〇])\s*(\S.*?)\s*$")
S_HEADING = re.compile(r"^\s*(?:S(?:CENE)?\s*#?|シーン)\s*[._-]?\s*(\d+[A-Za-z]?)(?:[\s　:：]+(.*?))?\s*$", re.I)
INT_HEADING = re.compile(r"^\s*(?:INT\.|EXT\.)\s*(\S.*?)\s*$", re.I)
NUM_HEADING = re.compile(r"^\s*(\d{1,3}[A-Za-z]?)[\s\t　]+(\S.{0,80}?)\s*$")
TRAILING_TIME = re.compile(rf"\s*(?:[（(]([^）)]*)[）)]|[／/・-]\s*({TIME_TOKEN})\s*)$", re.I)
SPEAKER = re.compile(r"^((?:[\w一-龥ぁ-んァ-ンー々〆]\s?){1,15}?)\s*(?:[（(][^）)]{1,20}[）)])?\s*(?:[「『]|[：:])")
# 仮名・漢字に挟まれた英字や記号（っ→] 、ー→. のようにフォントの対応表が壊れた字）。「……」の点の連続は含めない
SUSPECT = re.compile(rf"(?<=[{JP}])(?:[A-Za-z\u00c0-\u024f\u0400-\u04ff\]\[\}}\{{]+|\.)(?=[{JP}])")
PARTICLE_END = re.compile(r"[にをがはでとへもや]$")
NON_SHOOTING = re.compile(r"^(黒味|黒コマ|メインタイトル|サブタイトル|タイトル|エンドロール|エンドクレジット|アイキャッチ)$")
EPISODE = re.compile(r"(?:第\s*([0-9０-９一二三四五六七八九十]+)\s*話|#\s*([0-9０-９]+)|([0-9０-９]+)\s*話|EP(?:ISODE)?\.?\s*([0-9]+))", re.I)
KANJI_DIGITS = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}

def _kanji_number(text):
    text = unicodedata.normalize("NFKC", text)
    if text.isdigit(): return int(text)
    if text == "十": return 10
    if "十" in text:
        head, _, tail = text.partition("十")
        return (KANJI_DIGITS.get(head, 1) if head else 1) * 10 + (KANJI_DIGITS.get(tail, 0) if tail else 0)
    return KANJI_DIGITS.get(text)

def detect_episode(lines, name=None):
    # 表紙の行、表紙を1続きにした文字列（字間の空いた題字）、ファイル名の順に探す
    for line in lines[:40] + ["".join(lines[:200])] + ([name] if name else []):
        m = EPISODE.search(unicodedata.normalize("NFKC", line))
        if m:
            value = _kanji_number(next(g for g in m.groups() if g))
            if value: return str(value)
    return None

def _day_night(token):
    token = token.strip().upper()
    for code, words in TIME_WORDS.items():
        if token in (w.upper() for w in words): return code
    return None

def split_time(location):
    """見出しの末尾の時間帯（（夜）・／NIGHT 等）を外す。括弧に時間帯以外（回想・数日後）が混ざるときは残す"""
    text, day_night = location.strip(), None
    m = TRAILING_TIME.search(text)
    if m:
        inner = m.group(1) if m.group(1) is not None else m.group(2)
        parts = [p for p in re.split(r"[・、,/／\s]+", inner or "") if p]
        codes = [_day_night(p) for p in parts]
        found = [c for c in codes if c]
        if found:
            day_night = found[-1]
            rest = [p for p, c in zip(parts, codes) if not c]
            text = text[:m.start()].rstrip() + (f"（{'・'.join(rest)}）" if rest else "")
    return text.strip(" ・/／-"), day_night

def _heading(line, numbered_ok):
    text = line.strip()
    m = GLYPH_HEADING.match(text)
    if m: return None, m.group(1), "glyph"
    m = S_HEADING.match(text)
    if m: return m.group(1), m.group(2) or "", "s"
    m = INT_HEADING.match(text)
    if m: return None, m.group(1), "int"
    if numbered_ok:
        m = NUM_HEADING.match(text)
        if m and JP_CHAR.search(m.group(2)) and not m.group(2).startswith("「") and not re.search(r"[。、]", m.group(2)): return m.group(1), m.group(2), "number"
    return None

def _resolve_same(location, previous):
    # 「同・隣室」は直前の場所の「・」より前を引き継ぐ
    m = re.match(r"^同[・\s]*(.*)$", location)
    if not m or not previous: return location
    base = previous.split("・")[0].strip()
    return f"{base}・{m.group(1)}" if m.group(1) else base

def analyze_script(lines, layout=None, name=None):
    """行の並び（[{text,page,slot}] または文字列）から、シーン候補と読み取りの注意を作る"""
    if isinstance(lines, str):
        lines = [{"text": t, "page": None, "slot": None} for t in lines.replace("\r", "").split("\n")]
    layout = layout or {}
    texts = [l["text"] for l in lines]
    has_glyph = any(GLYPH_HEADING.match(t) or S_HEADING.match(t) or INT_HEADING.match(t) for t in texts)
    numbered_ok = not has_glyph
    heads = []
    last_number = 0
    for index, line in enumerate(lines):
        found = _heading(line["text"], numbered_ok)
        if not found: continue
        number, location, style = found
        if style == "number":
            n = int(re.match(r"\d+", number).group(0))
            if n <= last_number or n > last_number + 5: continue   # 番号が進まない行は見出しにしない（本文の数字）
            last_number = n
        heads.append((index, number, location, style))
    episode = detect_episode(texts[:heads[0][0]] if heads else texts, name)   # 話数は最初の見出しより前（表紙）とファイル名からだけ読む
    slots = layout.get("slotsPerPage")
    scenes, previous_location = [], None
    style_counts = {}
    for k, (index, number, raw_location, style) in enumerate(heads):
        style_counts[style] = style_counts.get(style, 0) + 1
        end = heads[k + 1][0] if k + 1 < len(heads) else len(lines)
        body = lines[index + 1:end]
        # 喫茶「灯台」のように括弧の閉じた見出しはそのまま。閉じない「や、」の後に続く「は台詞が同じ列に入ったとみる
        mixed = raw_location.count("「") != raw_location.count("」") or bool(re.search(r"」.*「", raw_location))
        location, day_night = split_time(raw_location.split("「")[0] if mixed else raw_location)
        location = _resolve_same(location, previous_location)
        previous_location = location or previous_location
        no = number or str(len(scenes) + 1)
        scene_no = f"{episode}-{int(no) if no.isdigit() else no}" if episode else f"S{no}"
        cast, synopsis, notes, suspects = [], "", [], []
        for l in body:
            clean = l["text"].strip()
            if not clean: continue
            speaker = SPEAKER.match(clean)
            name = re.sub(r"(?:[MＭМNＮ]|の声)$", "", re.sub(r"\s+", "", speaker.group(1))) if speaker else ""
            if name and not PARTICLE_END.search(name):
                if name not in cast: cast.append(name)
            elif not synopsis and not re.fullmatch(r"[×✕xX\s※]+", clean): synopsis = clean[:1000]
            if SUSPECT.search(clean): suspects.append(clean[:40])
        if not synopsis: synopsis = next((l["text"].strip()[:1000] for l in body if l["text"].strip()), "")
        if SUSPECT.search(raw_location): notes.append("場所の文字が化けている疑いがあります")
        if mixed: notes.append("見出しに台詞が混ざっている可能性があります（原本で場所を確かめてください）")
        if suspects: notes.append(f"文字化けの疑いがある行が{len(suspects)}行あります（例: {suspects[0]}）")
        if NON_SHOOTING.match(location): notes.append("撮影しない見出し（黒味・タイトル）の可能性があります")
        # 頁（1/8頁単位）: 見出しから次の見出しまでの行（縦書きは列）の幅を、1頁の行数で割る
        page_eighths, start = None, _position(lines[index], slots)
        if start is not None:
            finish = _position(lines[end], slots) if end < len(lines) else None
            if finish is None:
                last = next((_position(l, slots) for l in reversed(body) if l["text"].strip() and _position(l, slots) is not None), None)
                finish = last + 1 if last is not None else start + 1
            page_eighths = max(1, round((finish - start) / slots * 8))
        scenes.append({"sceneNo": scene_no, "location": location, "dayNight": day_night, "synopsis": synopsis, "cast": cast,
                       "estimatedMinutes": None, "pageEighths": page_eighths, "heading": raw_location.strip(), "reviewNotes": notes})
    suspect_total = sum(1 for t in texts if SUSPECT.search(t))
    warnings = []
    if suspect_total: warnings.append(f"文字化けの疑いがある行が{suspect_total}行あります。原本と見比べて直してください（自動では直しません）")
    if scenes and not any(s["dayNight"] for s in scenes): warnings.append("見出しに昼夜の記載がありません。香盤で昼夜を入れてください")
    if layout.get("orientation") == "vertical": warnings.append("縦書きのPDFを列ごとに組み直して読みました。改行の位置は原本と違うことがあります")
    meta = {"episode": episode, "headingStyles": style_counts, "orientation": layout.get("orientation"), "pages": layout.get("pages"),
            "slotsPerPage": slots, "suspectLines": suspect_total, "warnings": warnings}
    return scenes, meta

def _position(line, slots):
    if not slots or line.get("page") is None or line.get("slot") is None: return None
    return line["page"] * slots + line["slot"]

def script_scenes(text):
    return analyze_script(text)[0]

def table_from_text(text):
    lines = [line.strip() for line in text.replace("\r", "").split("\n") if line.strip()]
    if not lines: return []
    if any("\t" in line for line in lines): rows = [[cell.strip() for cell in line.split("\t")] for line in lines]
    elif any("," in line for line in lines): rows = list(csv.reader(lines))
    else: rows = [re.split(r"\s{2,}", line) for line in lines]
    width = max((len(row) for row in rows), default=0)
    return rows if width > 1 else []

def main():
    request = json.load(sys.stdin)
    data = base64.b64decode(request["base64"], validate=True)
    name = os.path.basename(str(request.get("name") or "upload"))
    ext = os.path.splitext(name)[1].lower()
    result = {"extractorName":"openingnight-local-python", "extractorVersion":VERSION,
              "rawSha256":hashlib.sha256(data).hexdigest(), "byteLength":len(data), "name":name}
    if ext == ".xlsx":
        result.update({"documentType":"workbook", "sheets":rows_from_xlsx(data), "text":""})
    elif ext == ".docx":
        text = text_from_docx(data); scenes, meta = analyze_script(text, name=name)
        result.update({"documentType":"text", "text":text, "scenes":scenes, "scriptMeta":meta})
    elif ext == ".pdf":
        parsed = pdf_script(data); text = parsed["text"]
        if not text.strip():
            result.update({"documentType":"pdf", "status":"ocr_pending", "unsupportedReason":"画像のみのPDFです。OCRは未対応です。シーンや売上行がないとは判定していません。", "text":""})
        else:
            scenes, meta = analyze_script(parsed["lines"], parsed["layout"], name)
            result.update({"documentType":"text", "text":text, "scenes":scenes, "scriptMeta":meta})
            rows = table_from_text(text)
            if rows: result["sheets"] = [{"name":"PDF text", "rows":rows}]
    elif ext in (".txt", ".csv"):
        try: text, encoding = data.decode("utf-8-sig"), "utf-8-sig"
        except UnicodeDecodeError: text, encoding = data.decode("cp932"), "cp932"
        result["sourceEncoding"] = encoding
        if ext == ".csv":
            rows = list(csv.reader(io.StringIO(text))); result.update({"documentType":"workbook", "sheets":[{"name":"CSV", "rows":rows}], "text":text})
        else:
            scenes, meta = analyze_script(text, name=name)
            result.update({"documentType":"text", "text":text, "scenes":scenes, "scriptMeta":meta})
    else: raise ValueError("対応形式は TXT、CSV、DOCX、XLSX、テキストPDFです")
    result.setdefault("status", "extracted")
    json.dump(result, sys.stdout, ensure_ascii=False)

if __name__ == '__main__':
    try: main()
    except Exception as exc:
        json.dump({"error":str(exc)}, sys.stdout, ensure_ascii=False); sys.exit(2)
