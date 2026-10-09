"""Private Container HTTP service for deterministic CSV/XLSX/PDF extraction."""
import base64,csv,datetime,hashlib,io,json,os,sys,zipfile
from http.server import BaseHTTPRequestHandler,ThreadingHTTPServer

MAX_BYTES=5*1024*1024
MAX_ROWS=10001
MAX_COLS=100

def scalar(value):
    if value is None:return ''
    if isinstance(value,datetime.datetime):return value.isoformat()
    if isinstance(value,datetime.date):return value.isoformat()
    if isinstance(value,float) and value.is_integer():return int(value)
    return value

def extract(request):
    name=os.path.basename(str(request.get('name','')));ext=os.path.splitext(name)[1].lower()
    raw=base64.b64decode(request.get('base64',''),validate=True)
    if not raw or len(raw)>MAX_BYTES:raise ValueError('ファイルは5MB以内にしてください')
    result={'name':name,'rawSha256':hashlib.sha256(raw).hexdigest(),'sourceHash':hashlib.sha256(raw).hexdigest(),'byteLength':len(raw),'extractorName':'openingnight-cloud-container','extractorVersion':'cloud-extract-3','status':'extracted'}
    if ext=='.csv':
        try:text=raw.decode('utf-8-sig');encoding='utf-8-sig'
        except UnicodeDecodeError:text=raw.decode('cp932');encoding='cp932'
        rows=[]
        for row in csv.reader(io.StringIO(text),strict=True):
            if len(rows)>=MAX_ROWS or len(row)>MAX_COLS:raise ValueError('見出し込み10001行・100列以内にしてください')
            rows.append(row)
        result.update(documentType='workbook',mediaType='text/csv',sourceEncoding=encoding,sheets=[{'name':'CSV','rows':rows,'formulaIssues':[]}])
    elif ext=='.xlsx':
        import openpyxl
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist())>40*1024*1024:raise ValueError('展開後のExcelが40MB上限を超えています')
        values=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
        formulas=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        sheets=[]
        try:
            if len(values.worksheets)>30:raise ValueError('シート数は30以内です')
            for sheet in values.worksheets:
                if (sheet.max_row or 0)>MAX_ROWS or (sheet.max_column or 0)>MAX_COLS:raise ValueError('1シートは見出し込み10001行・100列以内にしてください')
                rows=[[scalar(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
                issues=[]
                for row in formulas[sheet.title].iter_rows():
                    for cell in row:
                        if cell.data_type=='f' and (cell.row>len(rows) or cell.column>len(rows[cell.row-1]) or rows[cell.row-1][cell.column-1]==''):issues.append({'cell':cell.coordinate,'reason':'数式の保存済み計算値がありません。Excelで再計算・保存してください。'})
                sheets.append({'name':sheet.title,'rows':rows,'formulaIssues':issues})
        finally:values.close();formulas.close()
        result.update(documentType='workbook',mediaType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',sheets=sheets)
    elif ext=='.pdf':
        from extract_document import analyze_script,pdf_script,tables_from_pdf
        parsed=pdf_script(raw);text=parsed['text']
        if not text.strip():result.update(documentType='text',mediaType='application/pdf',text='',scenes=[],status='ocr_pending',unsupportedReason='画像PDFのOCRは未対応です。')
        else:
            sheets=tables_from_pdf(raw);scenes,meta=analyze_script(parsed['lines'],parsed['layout'],name)
            result.update(documentType='text',mediaType='application/pdf',text=text,scenes=scenes,scriptMeta=meta,sheets=sheets or [{'name':'PDF text','rows':[[line] for line in text.splitlines()[:MAX_ROWS]],'formulaIssues':[]}])
    elif ext in ('.txt','.docx'):
        from extract_document import analyze_script,text_from_docx
        if ext=='.docx':text=text_from_docx(raw);media='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        else:
            try:text=raw.decode('utf-8-sig');result['sourceEncoding']='utf-8-sig'
            except UnicodeDecodeError:text=raw.decode('cp932');result['sourceEncoding']='cp932'
            media='text/plain'
        scenes,meta=analyze_script(text,name=name)
        result.update(documentType='text',mediaType=media,text=text,scenes=scenes,scriptMeta=meta)
    else:raise ValueError('TXT、CSV、DOCX、XLSX、PDFだけを抽出できます')
    return result

class Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path!='/ready':self.send_error(404);return
        self.send_response(200);self.end_headers();self.wfile.write(b'ok')
    def do_POST(self):
        if self.path not in ('/extract','/analytics/build'):self.send_error(404);return
        try:
            length=int(self.headers.get('content-length','0'))
            limit=7*1024*1024 if self.path=='/extract' else 8*1024*1024
            if length<=0 or length>limit:self.send_error(413);return
            request=json.loads(self.rfile.read(length))
            if self.path=='/extract':output=extract(request)
            else:
                from cloud_analytics import build
                output=build(request)
            body=json.dumps(output,ensure_ascii=False,separators=(',',':')).encode('utf-8');status=200
            if len(body)>16*1024*1024:raise ValueError('Response exceeds 16MB')
        except Exception as error:body=json.dumps({'error':str(error)},ensure_ascii=False).encode('utf-8');status=422
        self.send_response(status);self.send_header('content-type','application/json; charset=utf-8');self.send_header('content-length',str(len(body)));self.end_headers();self.wfile.write(body)
    def log_message(self,format,*args):return

if __name__=='__main__':ThreadingHTTPServer(('0.0.0.0',8080),Handler).serve_forever()
